"""Fresh-browser acceptance of all four goods pages; temporary fixtures are removed in finally."""
import io
import json
import re
import socket
import subprocess
import sys
import time
import urllib.request
from pathlib import Path
import openpyxl
from playwright.sync_api import sync_playwright, expect
from verify_channel_inventory_ui import request, opener, BASE, LOGS
import verify_channel_inventory_ui as local_ui

BASE = 'http://127.0.0.1:8088'
local_ui.BASE = BASE

def login():
    captcha = request('/code')
    command = ['docker', 'exec', 'oms-local-redis', 'redis-cli', '--raw', 'GET', 'captcha_codes:' + captcha['uuid']]
    if sys.platform == 'win32':
        key = ('captcha_codes:' + captcha['uuid']).encode()
        with socket.create_connection(('127.0.0.1', 16379), timeout=10) as connection:
            connection.sendall(b'*2\r\n$3\r\nGET\r\n$' + str(len(key)).encode() + b'\r\n' + key + b'\r\n')
            reader = connection.makefile('rb')
            length = int(reader.readline()[1:].strip())
            assert length > 0, 'Fresh captcha not found'
            answer = reader.read(length).decode().strip('"')
    else:
        answer = subprocess.check_output(command, text=True, timeout=30).strip().strip('"')
    return request('/auth/login', {'username': 'admin', 'password': 'admin123', 'uuid': captcha['uuid'], 'code': answer})['data']['access_token']

def binary(path, token):
    req = urllib.request.Request(BASE + '/prod-api' + path, data=b'', headers={'Authorization': 'Bearer ' + token}, method='POST')
    with opener.open(req, timeout=30) as response:
        return response.read()

def main():
    LOGS.mkdir(exist_ok=True)
    token = login()
    print('Goods UI authenticated', flush=True)
    prefix = 'UIT' + str(int(time.time()))
    print('Temporary fixture scope:', prefix, flush=True)
    masters, products, batches = [], [], []
    try:
        def master(kind, data):
            ident = request('/goods/' + kind, data, token)['data']
            masters.append((kind, ident))
            return ident
        root = master('category', {'name': prefix, 'pid': 0})
        middle = master('category', {'name': '测试二级', 'pid': root})
        leaf = master('category', {'name': '测试三级', 'pid': middle})
        color = master('color', {'colorName': prefix + '-白', 'outColorCode': prefix + '-C'})
        size = master('size', {'sizeName': prefix + '-M', 'outSizeCode': prefix + '-S', 'sortOrder': 3})
        for i in range(12):
            data = {'skuSn': prefix + '-' + str(i), 'goodsSn': prefix, 'goodsName': prefix + '商品' + str(i), 'categoryCode': leaf, 'colorCode': color, 'sizeCode': size, 'marketPrice': 0, 'validity': 30}
            if i == 11:
                data.update(isGift=1, isFd=1, isPackage=1)
            products.append(request('/goods/info', data, token)['data'])
        first = request('/goods/info/list?pageNum=1&pageSize=5', {'goodsSn': prefix}, token)
        second = request('/goods/info/list?pageNum=2&pageSize=5', {'goodsSn': prefix}, token)
        assert first['total'] == second['total'] == 12 and len(first['rows']) == len(second['rows']) == 5
        assert {r['id'] for r in first['rows']}.isdisjoint({r['id'] for r in second['rows']})
        assert request('/goods/info/list', {'categoryCode': -999999}, token)['total'] == 0
        assert request('/goods/info/list', {'categoryCode': root, 'colorCode': color, 'sizeCode': size}, token)['total'] == 12
        print('API paging and combined master filters passed', flush=True)
        template = binary('/goods/goodsAdministration/importTemplate', token)
        def workbook(name, count, valid):
            book = openpyxl.load_workbook(io.BytesIO(template))
            sheet = book.worksheets[0]
            headers = [str(cell.value).split('(')[0] for cell in sheet[1]]
            for i in range(count):
                row = {'SKU': prefix + ('-IMP' if valid else '-BAD') + str(i), '货号': prefix, '商品名称': prefix + '导入商品' + str(i), '分类': prefix + '/测试二级/测试三级', '颜色': prefix + '-白' if valid else '尚未维护的测试颜色', '尺码': prefix + '-M', '市场价': 0, '有效期': 60, '商品描述': '浏览器验收', '是否福袋': 0, '是否赠品': 1, '是否套装': 1}
                sheet.append([row.get(header) for header in headers])
            path = LOGS / name
            book.save(path)
            return path
        bad_file = workbook('goods-ui-invalid.xlsx', 12, False)
        good_file = workbook('goods-ui-valid.xlsx', 2, True)
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(viewport={'width': 1440, 'height': 1000}, accept_downloads=True)
            context.add_cookies([{'name': 'Admin-Token', 'value': token, 'url': BASE}])
            page = context.new_page()
            page.set_default_timeout(20000)
            errors = []
            page.on('pageerror', lambda error: errors.append(str(error)))
            def goto(path):
                page.goto(BASE + path, wait_until='networkidle')
                assert '加载失败' not in page.locator('.app-container').inner_text()
            goto('/oms-goods/info?goodsSn=' + prefix)
            expect(page.locator('.goods-workspace .el-table__body-wrapper tr')).to_have_count(10)
            for label in ['赠品：是', '福袋：是', '套装：是', '赠品：否', '福袋：否', '套装：否']:
                assert label in page.locator('.app-container .el-table').inner_text()
            page.screenshot(path=str(LOGS / 'goods-info-desktop.png'), full_page=True)
            page.locator('.app-container .el-table').get_by_role('button', name='修改', exact=True).first.click()
            dialog = page.get_by_role('dialog', name='修改商品资料', exact=True)
            assert dialog.locator('.el-form-item').filter(has=page.locator('label').filter(has_text='SKU')).locator('input').is_disabled()
            assert dialog.locator('.el-select').count() == 3
            page.screenshot(path=str(LOGS / 'goods-editor.png'))
            dialog.get_by_role('button', name='取消', exact=True).click()
            page.locator('.app-container .el-table').get_by_role('button', name='详情', exact=True).first.click()
            page.locator('.goods-detail-drawer').get_by_text(prefix + ' / 测试二级 / 测试三级', exact=True).wait_for()
            page.screenshot(path=str(LOGS / 'goods-detail.png'))
            page.locator('.goods-detail-drawer .el-drawer__close-btn').click()
            with page.expect_download() as download:
                page.get_by_role('button', name='导出筛选结果').click()
            download.value.save_as(str(LOGS / 'goods-ui-export.xlsx'))
            assert openpyxl.load_workbook(LOGS / 'goods-ui-export.xlsx').active.max_row == 13
            print('Goods list, edit selections, detail and filtered export passed', flush=True)
            goto('/oms-goods/category')
            page.get_by_placeholder('搜索分类，保留上级路径').fill('测试三级')
            page.get_by_role('button', name='搜索').click()
            expect(page.locator('.app-container .el-table__body-wrapper tr')).to_have_count(3)
            assert prefix in page.locator('.app-container .el-table').inner_text()
            page.screenshot(path=str(LOGS / 'goods-category.png'), full_page=True)
            page.get_by_role('button', name='新增一级分类').click()
            page.get_by_role('dialog').get_by_text('无上级（一级分类）', exact=True).wait_for()
            page.get_by_role('dialog').get_by_role('button', name='取消', exact=True).click()
            for kind, label in [('color', '颜色'), ('size', '尺码')]:
                goto('/oms-goods/' + kind)
                page.get_by_placeholder('搜索' + label + '名称').fill(prefix)
                page.get_by_role('button', name='搜索').click()
                page.wait_for_load_state('networkidle')
                assert '12 个' in page.locator('.app-container .el-table').inner_text()
                page.screenshot(path=str(LOGS / ('goods-' + kind + '.png')), full_page=True)
                page.locator('.app-container .el-table').get_by_role('button', name='查看', exact=True).first.click()
                page.wait_for_url(lambda url: '/oms-goods/info' in url and kind + 'Code=' in url)
                page.wait_for_load_state('networkidle')
                expect(page.locator('.goods-workspace .el-table__body-wrapper tr')).to_have_count(10)
            print('Category tree and color/size related-goods navigation passed', flush=True)
            goto('/oms-goods/info?goodsSn=' + prefix)
            page.get_by_role('button', name='新增商品').click()
            add = page.get_by_role('dialog', name='新增商品资料', exact=True)
            def field(label):
                return add.locator('.el-form-item').filter(has=page.locator('label').filter(has_text=re.compile('^' + re.escape(label) + '$'))).locator('input').first
            field('SKU').fill(prefix + '-MANUAL')
            field('货号').fill(prefix)
            field('商品名称').fill(prefix + '手工新增')
            for index, label in enumerate([prefix + ' / 测试二级 / 测试三级', prefix + '-白', prefix + '-M']):
                add.locator('.el-select').nth(index).locator('input').click()
                add.locator('.el-select').nth(index).locator('input').fill(label)
                page.locator('.el-select-dropdown:visible .el-select-dropdown__item').filter(has_text=label).first.click()
            field('市场价').fill('0')
            field('商品名称').click()
            with page.expect_response(lambda response: response.url.endswith('/goods/info') and response.request.method == 'POST') as created:
                add.get_by_role('button', name='保存商品', exact=True).click()
            saved = created.value.json()
            assert saved['code'] == 200, saved.get('msg')
            products.append(saved['data'])
            add.wait_for(state='hidden')
            print('Manual creation through all three master selections passed', flush=True)
            def upload(path):
                page.get_by_role('button', name='导入商品').click()
                dialog = page.get_by_role('dialog', name='导入商品', exact=True)
                dialog.locator('input[type=file]').set_input_files(str(path))
                with page.expect_response(lambda response: '/goodsAdministration/import' in response.url and response.request.method == 'POST') as result:
                    dialog.get_by_role('button', name='上传并校验').click()
                body = result.value.json()
                assert body['code'] == 200, body.get('msg')
                batches.append(body['msg'])
                preview = page.get_by_role('dialog', name='商品导入校验预览', exact=True)
                preview.get_by_text('共 ' + ('12' if path == bad_file else '2') + ' 行', exact=False).wait_for()
                return preview
            preview = upload(bad_file)
            assert preview.get_by_role('button', name='确认导入 12 行商品').is_disabled()
            preview.locator('.el-pagination .btn-next').click()
            expect(preview.locator('.el-table__body-wrapper tr')).to_have_count(2)
            assert '13' in preview.locator('.el-table').inner_text()
            with page.expect_download() as download:
                preview.get_by_role('button', name='下载错误明细').click()
            download.value.save_as(str(LOGS / 'goods-ui-errors.xlsx'))
            assert openpyxl.load_workbook(LOGS / 'goods-ui-errors.xlsx').active.max_row == 13
            page.screenshot(path=str(LOGS / 'goods-import-errors.png'))
            preview.get_by_role('button', name='关闭', exact=True).click()
            preview = upload(good_file)
            preview.get_by_role('button', name='确认导入 2 行商品').click()
            page.get_by_role('dialog', name='导入完成', exact=True).get_by_text('已成功新增 2 个商品。', exact=True).wait_for()
            page.get_by_role('dialog', name='导入完成', exact=True).get_by_role('button', name='确定', exact=True).click()
            imported = request('/goods/info/list?pageSize=100', {'goodsName': prefix + '导入商品'}, token)['rows']
            products.extend(r['id'] for r in imported)
            assert len(imported) == 2 and all(r['validity'] == '60' and r['isGift'] == 1 and r['isPackage'] == 1 for r in imported)
            request('/goods/goodsAdministration/toExamine?import_batch=' + batches[-1], {}, token)
            assert request('/goods/info/list', {'goodsName': prefix + '导入商品'}, token)['total'] == 2
            print('Invalid import paging/errors and valid idempotent confirmation passed', flush=True)
            page.set_viewport_size({'width': 390, 'height': 844})
            page.wait_for_timeout(400)
            page.screenshot(path=str(LOGS / 'goods-info-mobile.png'), full_page=True)
            page.locator('.app-container .el-table').get_by_role('button', name='详情', exact=True).first.click()
            page.wait_for_timeout(400)
            assert page.locator('.goods-detail-drawer').bounding_box()['x'] < 1
            page.screenshot(path=str(LOGS / 'goods-detail-mobile.png'))
            assert not errors, errors
            context.close()
            browser.close()
            print('PASS four goods pages, import, export, related navigation, desktop/mobile; no browser errors', flush=True)
    finally:
        # Discover any product created before an assertion failed, then delete only this run's fixture IDs.
        try:
            found = request('/goods/info/list?pageSize=100', {'goodsSn': prefix}, token)['rows']
            all_ids = sorted(set(products + [r['id'] for r in found]))
            if all_ids:
                request('/goods/info/' + ','.join(map(str, all_ids)), token=token, method='DELETE')
            for kind, ident in reversed(masters):
                request('/goods/' + kind + '/' + str(ident), token=token, method='DELETE')
            # Preview audit rows are scoped to exact UUIDs returned by this run.
            if batches:
                payload = json.dumps(batches)
                cleanup = 'import sys,json; sys.path.insert(0,"docker/local"); from local_db import execute,quote; ids=json.loads(sys.argv[1]); values=",".join(map(quote,ids)); execute("DELETE FROM goods_import_row WHERE import_batch IN ("+values+"); DELETE FROM goods_import_batch WHERE import_batch IN ("+values+")", "qm_oms_saas_commodity")'
                command = ['python3', '-c', cleanup, payload]
                if sys.platform == 'win32':
                    command = ['wsl', '-d', 'Ubuntu-24.04', '--cd', str(Path(__file__).resolve().parents[2]).replace('D:\\', '/mnt/d/').replace('\\', '/'), '--'] + command
                subprocess.run(command, check=True, timeout=30)
            print('Temporary goods UI fixtures removed', flush=True)
        finally:
            request('/auth/logout', {}, token, method='DELETE')

if __name__ == '__main__':
    main()
