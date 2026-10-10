"""Browser and API purchase acceptance. All stock uses unique new SKUs and is cleaned in finally."""
import io
import json
import time
import subprocess
import urllib.request
import urllib.error
import openpyxl
from playwright.sync_api import sync_playwright, expect
from verify_goods_workspace_ui import login, request, BASE, LOGS, opener

ROOT='/supplychain/purchaseWorkspace'

def main():
    token=login();prefix='PURUI'+str(int(time.time()))
    print('Temporary purchase fixture scope:', prefix, flush=True)
    def post(path, body=None): return request(ROOT+path,{} if body is None else body,token)
    def read(kind, key): return request(f'{ROOT}/{kind}/{key}',token=token)['data']
    def rejected(path, body=None, method=None):
        req=urllib.request.Request(BASE+'/prod-api'+path,data=None if body is None else json.dumps(body).encode(),headers={'Authorization':'Bearer '+token,'Content-Type':'application/json'},method=method)
        try: response=opener.open(req,timeout=30)
        except urllib.error.HTTPError as error: response=error
        with response: result=json.load(response)
        assert result.get('code',response.status)!=200,(path,result)
    def master(kind, body): return request('/supplychain/'+kind,body,token)['data']
    try:
        owner=master('owner',{'ownerCode':prefix+'-O','ownerName':prefix+'货主','isEnable':2})
        warehouse=master('realStore',{'realStoreCode':prefix+'-W','wmsName':prefix+'实仓','status':2,'wmsType':1})
        relation=master('ownerWarehouse',{'ownerId':owner,'realStoreId':warehouse,'status':2})
        master('simulationStore',{'wmsSimulationCode':prefix+'-V','wmsSimulationName':prefix+'虚仓','inboundMode':2,'outboundMode':2,'ownerWarehouseId':relation,'status':2})
        # The existing catalog only supplies dimension IDs, never an existing stock SKU.
        catalog=request('/goods/info/options',token=token)['data']
        product={'skuSn':prefix+'-SKU','goodsSn':prefix,'goodsName':prefix+'验收商品','marketPrice':10,'validity':30}
        for key, collection in [('categoryCode','category'),('colorCode','color'),('sizeCode','size')]:
            product[key]=catalog[collection][0]['id']
        product['categoryCode']=next(c['id'] for c in catalog['category'] if c['level']==3)
        request('/goods/info',product,token)
        template_request=urllib.request.Request(BASE+'/prod-api'+ROOT+'/importTemplate',data=b'',headers={'Authorization':'Bearer '+token},method='POST')
        with opener.open(template_request) as response: template_bytes=response.read()
        def preview(import_rows):
            wb=openpyxl.load_workbook(io.BytesIO(template_bytes))
            for row in import_rows: wb.active.append(row)
            output=io.BytesIO();wb.save(output)
            boundary='purchase-ui-boundary'
            payload=(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="purchase.xlsx"\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n'.encode()+output.getvalue()+f'\r\n--{boundary}--\r\n'.encode())
            req=urllib.request.Request(BASE+'/prod-api'+ROOT+'/importPreview',data=payload,headers={'Authorization':'Bearer '+token,'Content-Type':'multipart/form-data; boundary='+boundary})
            with opener.open(req,timeout=30) as response:return json.load(response)
        assert preview([[product['skuSn'],2.5,10]])['code']==200
        assert preview([[product['skuSn'],2.5,10],[product['skuSn'],3,2]])['code']!=200
        assert preview([[product['skuSn'],2.5,-1]])['code']!=200
        assert request(ROOT+'/purchase/list?keyword='+prefix,token=token)['total']==0,'Import preview must not create documents'
        supplier=post('/supplier',{'supplierName':prefix+'供应商','companyName':prefix+'有限公司','contactUser':'验收联系人','contactTel':'020-12345678','status':2})['data']
        supplier_sn=read('supplier',supplier)['supplierSn']
        with sync_playwright() as p:
            browser=p.chromium.launch(headless=True);context=browser.new_context(viewport={'width':1536,'height':1000},accept_downloads=True)
            context.add_cookies([{'name':'Admin-Token','value':token,'url':BASE}]);page=context.new_page();page.set_default_timeout(20000)
            errors=[];page.on('pageerror',lambda error:errors.append(str(error)))
            page.on('response', lambda r: print('Purchase save response:',json.dumps(r.json(),ensure_ascii=True),flush=True) if r.url.endswith('/purchaseWorkspace/purchase') and r.request.method=='POST' else None)
            page.goto(BASE+'/oms-supplychain/poInfo?create=1&supplierSn='+supplier_sn,wait_until='networkidle')
            dialog=page.locator('.el-dialog:visible');expect(dialog).to_be_visible()
            def field(label): return dialog.locator('.el-form-item').filter(has=page.locator('label').filter(has_text=label)).first
            field('采购名称').locator('input').fill(prefix+'采购')
            field('收货虚仓').locator('.el-select').click()
            page.locator('.el-select-dropdown:visible .el-select-dropdown__item').filter(has_text=prefix+'虚仓').click()
            dialog.get_by_role('button',name='手工添加 SKU').click()
            row=dialog.locator('.el-table__body-wrapper tr').first
            row.locator('input').nth(0).fill(product['skuSn'])
            row.locator('input').nth(1).fill('10');row.locator('input').nth(2).fill('2.50');row.locator('input').nth(2).blur()
            expect(dialog).to_contain_text('25.00')
            dialog.get_by_role('button',name='保存草稿',exact=True).click();expect(dialog).not_to_be_visible()
            purchase=request(ROOT+'/purchase/list?supplierSn='+supplier_sn,token=token)['rows'][0];pid=purchase['id']
            assert float(purchase['moneyExpected'])==25 and purchase['numberExpected']==10
            page.goto(BASE+'/oms-supplychain/poInfo?poSn='+purchase['poSn'],wait_until='networkidle')
            page.get_by_role('button',name='详情',exact=True).last.click()
            drawer=page.locator('.el-drawer:visible');expect(drawer).to_contain_text(product['skuSn'])
            assert read('purchase',pid)['poState']==1,'Opening details must not approve'
            drawer.get_by_role('button',name='审核采购',exact=True).click();page.locator('.el-message-box').get_by_role('button',name='确定',exact=True).click()
            expect(drawer.get_by_role('button',name='创建到货单',exact=True)).to_be_visible()
            drawer.get_by_role('button',name='创建到货单',exact=True).click()
            quantities=page.locator('.el-dialog:visible');quantities.locator('.el-input-number input').fill('6')
            quantities.get_by_placeholder('选择本批次到货日期').fill('2026-10-20');quantities.get_by_placeholder('填写送货约定等信息').fill('分批到货验收')
            quantities.get_by_role('button',name='创建待审核到货单',exact=True).click();expect(quantities).not_to_be_visible()
            expect(drawer).to_contain_text('待审核')
            n=read('purchase',pid)['receipts'][0];nid=n['id']
            assert n['remarks']=='分批到货验收' and str(n['expectedCallbackTime']).startswith('2026-10-20')
            assert read('purchase',pid)['receiptCount']==1 and not read('purchase',pid)['editable']
            drawer.locator('.el-tab-pane:visible').get_by_role('button',name='查看并审核',exact=True).last.click()
            expect(drawer.get_by_role('button',name='审核到货单',exact=True)).to_be_visible()
            assert read('receipt',nid)['noState']==2,'Viewing receipt must not approve'
            drawer.get_by_role('button',name='审核到货单',exact=True).click()
            expect(page.locator('.el-message-box')).to_contain_text('自动收货并增加库存')
            with page.expect_response(lambda r: '/receipt/'+str(nid)+'/approve' in r.url and r.request.method=='POST') as approved:
                page.locator('.el-message-box').get_by_role('button',name='确定',exact=True).click()
            assert approved.value.json()['code']==200,approved.value.json()
            t=read('receipt',nid)['tickets'][0];tid=t['id']
            page.goto(BASE+'/oms-supplychain/wmsTickets?sn='+t['sn'],wait_until='networkidle')
            page.get_by_role('button',name='详情',exact=True).last.click();drawer=page.locator('.el-drawer:visible')
            expect(drawer).to_contain_text(product['skuSn'])
            expect(drawer).to_contain_text('全部实收入账')
            assert read('receipt',nid)['noState']==4 and read('purchase',pid)['numberActually']==6
            assert t['statusTicket']==2 and t['inventoryStatus']==2
            expect(drawer.get_by_role('button',name='重试虚拟入库',exact=True)).to_have_count(0)
            page.screenshot(path=str(LOGS/'purchase-ui-ticket-detail.png'),full_page=True)
            post(f'/ticket/{tid}/post');assert read('purchase',pid)['numberActually']==6
            rejected(ROOT+f'/ticket/{tid}/cancel',{'reason':'不允许作废已入账单'})
            rejected('/supplychain/tickets/'+str(tid),method='DELETE')
            rejected('/supplychain/tmp/submitExamine?no_sn='+n['noSn'])
            rejected(ROOT+'/purchase/list?company_code=OTHER')
            line=read('purchase',pid)['lines'][0]
            rejected(ROOT+f'/purchase/{pid}/receipt',{'lines':[{'lineId':line['id'],'quantity':5}]})
            n2=post(f'/purchase/{pid}/receipt',{'lines':[{'lineId':line['id'],'quantity':4}]})['data']
            t2=post(f'/receipt/{n2}/approve')['data'];l2=read('ticket',t2)['lines'][0]
            assert read('receipt',n2)['noState']==4
            post(f'/ticket/{t2}/completeVirtual');post(f'/ticket/{t2}/post')
            assert read('purchase',pid)['poState']==4 and float(read('purchase',pid)['moneyActually'])==25
            for slug in ['supplier','poInfo','wmsTickets']:
                page.goto(BASE+'/oms-supplychain/'+slug,wait_until='networkidle');page.screenshot(path=str(LOGS/('purchase-ui-'+slug+'.png')),full_page=True)
            page.get_by_text('商品明细',exact=True).click();expect(page.locator('.purchase-workspace .el-table')).to_contain_text(product['skuSn'])
            expect(page.get_by_role('button',name='商品明细',exact=True)).to_have_attribute('aria-pressed','true')
            headers=page.locator('.purchase-workspace > .el-table .el-table__header-wrapper').first.inner_text()
            assert headers.count('库存入账')==1 and headers.count('计划 / 实收')==1,headers
            page.screenshot(path=str(LOGS/'purchase-ui-lines.png'),full_page=True)
            with page.expect_download() as dl: page.get_by_role('button',name='导出筛选结果').click()
            download=dl.value;target=LOGS/'purchase-lines-export.xlsx';download.save_as(target)
            wb=openpyxl.load_workbook(target);assert wb.active.max_row>=3
            assert not errors,errors
            rejected('/goods/api/goods/purchaseLookup?company_code=QM',{'skuSn':product['skuSn']})
            browser.close()
        print('PASS browser purchase creation, explicit approvals, automatic virtual receipt posting, split receipts, retry, guards, links, detail export',flush=True)
    finally:
        subprocess.run(['wsl','-d','Ubuntu-24.04','--','python3','/mnt/d/Users/yueguang/OmsSaaSCloud/docker/local/cleanup_purchase_ui.py',prefix],check=True)
        request('/auth/logout',{},token,method='DELETE')

if __name__=='__main__':main()
